import json
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


path = Path(__file__).with_name("final.docx")
doc = Document(path)
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def outline_level(paragraph):
    ppr = paragraph.style.element.find("w:pPr", namespaces=ns)
    if ppr is None:
        return None
    level = ppr.find("w:outlineLvl", namespaces=ns)
    return level.get(f"{{{ns['w']}}}val") if level is not None else None


headings = [
    paragraph.text.strip()
    for paragraph in doc.paragraphs
    if outline_level(paragraph) is not None
]
toc = [
    paragraph.text.strip()
    for paragraph in doc.paragraphs
    if paragraph.style.name.lower().startswith("toc")
]
all_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        all_text += "\n" + "\t".join(cell.text for cell in row.cells)

expected_headings = [
    "10.1 实际运行与数据来源",
    "10.2 默认参数检索结果分析",
    "10.3 K-mer 索引耗时与哈希负载对比实验",
    "10.4 边界情况与异常处理",
]
chapter_10_headings = [
    heading for heading in headings if heading.startswith("10.")
]

expected_phrases = [
    "K=2 平均每条 Read 产生 188.96 个候选",
    "累计 9448 个候选",
    "累计耗时降至 35.00 μs",
    "成功数下降为 37",
    "相对偏移 6～9",
    "K=10 已命中的 37 条 Read 与 K=6 的匹配位置完全一致",
    "Read编号、序列、匹配排名、匹配位置、汉明距离",
]
unexpected_phrases = [
    "10.1 实际运行与数据来源：待补充",
    "10.2 默认参数检索结果分析：待补充",
    "10.3 K-mer 索引耗时与哈希负载对比实验：待补充",
]

table_6 = [[cell.text for cell in row.cells] for row in doc.tables[5].rows]
table_7 = [[cell.text for cell in row.cells] for row in doc.tables[6].rows]

with zipfile.ZipFile(path) as archive:
    xml_parts = {
        name: archive.read(name)
        for name in archive.namelist()
        if name.endswith(".xml")
    }
    document_root = etree.fromstring(xml_parts["word/document.xml"])
    vertical_tabs = sum(part.count(b"\x0b") for part in xml_parts.values())
    soft_breaks = len(
        document_root.xpath(
            ".//w:br[not(@w:type) or @w:type='textWrapping']",
            namespaces=ns,
        )
    )
    fields = document_root.xpath(".//w:instrText/text()", namespaces=ns)

errors = []
if chapter_10_headings != expected_headings:
    errors.append("Chapter 10 headings do not match.")
if not all(phrase in all_text for phrase in expected_phrases):
    errors.append("One or more experiment statements are missing.")
if any(phrase in all_text or any(phrase in item for item in toc)
       for phrase in unexpected_phrases):
    errors.append("Stale pending experiment headings remain.")
if table_7[1:] != [
    ["2", "2999", "16", "156", "0.002", "9448 / 465.70 μs / 50"],
    ["6", "2995", "2128", "463", "0.260", "84 / 35.00 μs / 50"],
    ["10", "2991", "2987", "650", "0.365", "37 / 33.60 μs / 37"],
]:
    errors.append("Table 7 metrics do not match source data.")
if vertical_tabs or soft_breaks:
    errors.append("Forbidden line-break characters remain.")
if not any("TOC" in field for field in fields):
    errors.append("TOC field is missing.")

print(json.dumps(
    {
        "errors": errors,
        "chapter_10_headings": chapter_10_headings,
        "toc_chapter_10": [item for item in toc if item.startswith("10.")],
        "table_6": table_6,
        "table_7": table_7,
        "vertical_tabs": vertical_tabs,
        "soft_breaks": soft_breaks,
        "passed": not errors,
    },
    ensure_ascii=False,
    indent=2,
))
raise SystemExit(0 if not errors else 1)
