from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\Backup\Documents\GitHub\Tongji-CS-2026-spring-my-archive\数据结构课程设计")
SOURCE = ROOT / "tmp" / "report_revision" / "source.docx"
OUTPUT = ROOT / "tmp" / "report_revision" / "structured.docx"


def set_run_font(run, size=None, bold=None, superscript=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if superscript is not None:
        run.font.superscript = superscript


def configure_style(style, size, bold=False, outline_level=None):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "宋体")
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if outline_level is not None:
        ppr = style.element.get_or_add_pPr()
        outline = ppr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            ppr.append(outline)
        outline.set(qn("w:val"), str(outline_level))


def set_heading_style(paragraph, level):
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.first_line_indent = Pt(0)
    for run in paragraph.runs:
        set_run_font(run, {1: 16, 2: 12, 3: 14}[level], bold=True)


def add_body(doc, text):
    p = doc.add_paragraph(style="报告正文")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, 10.5)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(text)
    set_heading_style(p, level)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="报告表题")
    p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        set_run_font(run, 10.5)
    return p


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def style_table(table, widths):
    table.autofit = False
    table.alignment = 0
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "A7BAC4")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            set_repeat_header(row)
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)
        for column_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[column_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                tc_pr = cell._tc.get_or_add_tcPr()
                shading = tc_pr.find(qn("w:shd"))
                if shading is None:
                    shading = OxmlElement("w:shd")
                    tc_pr.append(shading)
                shading.set(qn("w:fill"), "D9EAF2")
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, 10.5, bold=row_index == 0)


def add_table(doc, rows, widths):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    style_table(table, widths)
    return table


def remove_body_range(start_element, end_element):
    parent = start_element.getparent()
    children = list(parent)
    start = children.index(start_element)
    end = children.index(end_element)
    for element in children[start:end]:
        parent.remove(element)


def move_before(element, anchor):
    anchor.addprevious(element)


doc = Document(SOURCE)

for heading_name in ("Heading 1", "Heading 2", "Heading 3"):
    if heading_name not in [style.name for style in doc.styles]:
        doc.styles.add_style(heading_name, WD_STYLE_TYPE.PARAGRAPH)

configure_style(doc.styles["Heading 1"], 16, bold=True, outline_level=0)
configure_style(doc.styles["Heading 2"], 12, bold=True, outline_level=1)
configure_style(doc.styles["Heading 3"], 14, bold=True, outline_level=2)
for name in ("Heading 1", "Heading 2", "Heading 3"):
    style = doc.styles[name]
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.keep_together = True
    style.paragraph_format.first_line_indent = Pt(0)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
doc.styles["Heading 3"].paragraph_format.space_before = Pt(8)
doc.styles["Heading 3"].paragraph_format.space_after = Pt(4)

if "报告目录标题" not in [style.name for style in doc.styles]:
    toc_title_style = doc.styles.add_style("报告目录标题", WD_STYLE_TYPE.PARAGRAPH)
else:
    toc_title_style = doc.styles["报告目录标题"]
configure_style(toc_title_style, 16, bold=True)
toc_title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
toc_title_style.paragraph_format.space_before = Pt(0)
toc_title_style.paragraph_format.space_after = Pt(12)
toc_title_style.paragraph_format.first_line_indent = Pt(0)

for style_name, left_indent in (("TOC 1", 0), ("TOC 2", 14), ("TOC 3", 28)):
    if style_name in [style.name for style in doc.styles]:
        style = doc.styles[style_name]
        configure_style(style, 10.5, bold=False)
        style.paragraph_format.left_indent = Pt(left_indent)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.15

# Convert report headings to real Word heading styles so the automatic TOC is reliable.
for paragraph in doc.paragraphs:
    if paragraph.style.name == "报告一级标题":
        set_heading_style(paragraph, 1)
    elif paragraph.style.name == "报告二级标题":
        set_heading_style(paragraph, 2)
    elif paragraph.style.name == "报告三级标题":
        set_heading_style(paragraph, 3)

chapter_10 = next(p for p in doc.paragraphs if p.text.strip().startswith("10 系统运行结果"))
references = next(p for p in doc.paragraphs if p.text.strip() == "参考文献")
remove_body_range(chapter_10._p, references._p)

new_blocks = []


def capture(element):
    new_blocks.append(element)
    return element


capture(add_heading(doc, "10 系统运行结果及分析", 1)._p)
capture(add_heading(doc, "10.1 实际运行与数据来源：待补充", 2)._p)
capture(
    add_body(
        doc,
        "待补充。实验应直接运行发布目录中的 DNA_SearchEngine.exe，依次导入 "
        "test_reference.fasta 与 test_reads.txt，在 K=6、允许错配数 k=2 的条件下显式建立索引、"
        "执行 50 条 Reads 的批量检索并导出 CSV。下表中的运行指标必须取自同一次实际操作，"
        "不得混用旧 CSV、随机生成数据或不同 K 值的日志。"
    )._p
)
capture(add_caption(doc, "表 6  默认参数实际运行记录（待补充）")._p)
capture(
    add_table(
        doc,
        [
            ["记录项", "实际值", "填写依据"],
            ["可执行程序", "DNA_SearchEngine.exe；构建时间待填", "执行程序目录中的实际文件"],
            ["Reference", "test_reference.fasta；3000 bp", "导入完成日志与运行统计"],
            ["Reads", "test_reads.txt；50 条，每条 30 bp", "Reads 列表与导入完成日志"],
            ["检索参数", "K=6，k=2", "参数设置区"],
            ["索引指标", "不同 K-mer、耗时、哈希负载：待填", "建立索引后的运行统计与日志"],
            ["批量检索", "成功数、失败数、候选总数、耗时：待填", "检索完成日志"],
            ["CSV 导出", "文件名、记录行数：待填", "导出文件与程序提示"],
        ],
        [1800, 2800, 3700],
    )._tbl
)

capture(add_heading(doc, "10.2 默认参数检索结果分析：待补充", 2)._p)
capture(
    add_body(
        doc,
        "待补充。完成默认参数实验后，应依据本次导出的 CSV 统计成功 Read 数、未命中 Read 数、"
        "汉明距离为 0、1、2 的结果数量、候选位置总数和检索耗时，并抽取若干记录重新对照 "
        "Reference。核对内容包括：Read 前 K 位种子是否对应候选位置，CSV 记录起点上的完整 "
        "Read 汉明距离是否与导出值一致，以及错配偏移和碱基替换是否与界面精确比对详情一致。"
    )._p
)
capture(
    add_body(
        doc,
        "结果分析应区分“候选位置”和“成功匹配位置”：种子命中只负责缩小范围，完整 Read 仍需"
        "逐字符验证；一个 Read 若存在多个合法位置，CSV 应保留全部位置，而最佳位置按汉明距离"
        "升序、起点升序确定。所有具体数字均在实验完成后填写。"
    )._p
)

capture(add_heading(doc, "10.3 K-mer 索引耗时与哈希负载对比实验：待补充", 2)._p)
capture(
    add_body(
        doc,
        "待补充。保持同一份 3000 bp Reference、同一份 50 条 Reads、允许错配数 k=2 和相同"
        "运行环境不变，依次设置 K=2、6、10。每次修改 K 后确认旧索引失效，再点击“建立索引”，"
        "记录界面显示的不同 K-mer 数、建索引耗时和哈希负载；随后执行批量检索，记录候选总数"
        "和检索耗时。"
    )._p
)
capture(add_caption(doc, "表 7  不同 K 值的索引与检索对比（待补充）")._p)
capture(
    add_table(
        doc,
        [
            ["K", "索引位置数", "不同 K-mer", "建索引耗时/μs", "哈希负载", "候选总数/检索耗时"],
            ["2", "2999", "待填", "待填", "待填", "待填"],
            ["6", "2995", "待填", "待填", "待填", "待填"],
            ["10", "2991", "待填", "待填", "待填", "待填"],
        ],
        [700, 1200, 1500, 1550, 1200, 2150],
    )._tbl
)
capture(
    add_body(
        doc,
        "实验完成后，应结合实际数据分析 K 增大时不同 K-mer 数、负载和候选规模的变化。索引耗时"
        "会受到系统调度和计时粒度影响，正文只解释本次实测结果，不以单次微秒差异推导普遍结论。"
    )._p
)

capture(add_heading(doc, "10.4 边界情况与异常处理", 2)._p)
capture(
    add_body(
        doc,
        "系统围绕课程题目的主要边界进行了针对性核对。下表只保留与算法正确性和用户操作直接"
        "相关的场景，不展开工程化压力测试或外部内存检测工具结果。"
    )._p
)
capture(add_caption(doc, "表 8  典型边界与异常场景")._p)
capture(
    add_table(
        doc,
        [
            ["测试场景", "预期结果", "实际表现"],
            [
                "完全匹配、1 处和 2 处错配",
                "k=0 仅接受完全匹配；k=1 接受至多 1 处；k=2 接受至多 2 处",
                "符合阈值定义，错配位置独立记录",
            ],
            [
                "Reference 首尾合法起点",
                "起点 0 与 N-ReadLength 均可检索，末字符参与比较",
                "均可返回，不越过 Reference 边界",
            ],
            [
                "重复 K-mer 与多候选位置",
                "同一键保存全部出现位置，逐个验证并保留全部合法结果",
                "候选不覆盖，结果按距离和起点稳定排序",
            ],
            [
                "空文件、非法碱基或格式错误",
                "拒绝导入并给出明确原因，不保留半加载数据",
                "通过 Snackbar、状态栏或文件错误信息提示",
            ],
            [
                "缺少 Reference、Reads 或有效索引",
                "阻止对应操作并提示缺失条件",
                "检索入口不会在状态不完整时继续执行",
            ],
            [
                "修改 K",
                "旧索引自动失效；再次建立索引时按新 K 重建",
                "旧结果清除，未重建前检索会提示索引无效",
            ],
            [
                "重新导入或生成数据",
                "清除旧检索结果和当前选中状态",
                "Reference 更新时同时清索引；Reads 更新时保留有效索引",
            ],
        ],
        [1850, 3300, 3150],
    )._tbl
)

capture(add_heading(doc, "11 系统安装、运行与操作说明", 1)._p)
capture(add_heading(doc, "11.1 直接运行", 2)._p)
capture(
    add_body(
        doc,
        "进入“250021_张三_计算机技术_执行程序”目录，双击 DNA_SearchEngine.exe。必须保留 "
        "EXE、Qt DLL、MinGW DLL、platforms、imageformats、styles 等插件目录的相对位置，"
        "不能只复制 EXE。程序不依赖 Qt Creator，正常发布包无需重新编译。"
    )._p
)
capture(add_heading(doc, "11.2 推荐操作步骤", 2)._p)
capture(
    add_body(
        doc,
        "第一步，点击“生成参考基因组”或导入合法 Reference。第二步，点击“随机生成 Reads”或"
        "导入 Reads。第三步，设置 K 和允许错配数并点击“建立索引”。第四步，在 Reads 列表中"
        "选择一条 Read 执行单条检索，或点击“高通量批量检索”。第五步，在全局概览、局部沙盘、"
        "检索结果和候选位置表中选择结果。第六步，按需点击“导出结果”保存 CSV。"
    )._p
)
capture(
    add_body(
        doc,
        "若改变 K，必须重新点击“建立索引”；改变允许错配数只需重新检索。相同 Reference 允许"
        "重复建立索引，每次都会刷新构建时间和哈希指标。清空检索结果不会删除 Reference、"
        "Reads 或有效索引。"
    )._p
)
capture(add_heading(doc, "11.3 输入文件格式", 2)._p)
capture(
    add_body(
        doc,
        "Reference 文件可以是 FASTA，也可以是只含 A/T/C/G 的纯文本，多行会拼接为一条序列。"
        "Reads 文件可以是 FASTA，或每行一条 Read 的纯文本。随发布材料提供 "
        "test_reference.fasta 和 test_reads.txt，用于验证导入、建索引、单条检索和批量检索；"
        "使用它们时推荐保持 K=6、k=2、Read 长度 30。"
    )._p
)
capture(add_heading(doc, "11.4 从源码构建", 2)._p)
capture(
    add_body(
        doc,
        "在已安装 Qt 6.11.1 MinGW 64-bit、匹配的 GCC、CMake 和 Ninja 的环境中，可先运行 "
        "tools/configure-release.cmd，再运行 tools/build-release.cmd。构建成功后运行 "
        "tools/deploy-release.cmd，由 windeployqt 复制依赖。若 Qt 未安装在 D:\\Qt，应先设置 "
        "QT_PREFIX、DNA_MINGW_BIN、DNA_CMAKE 和 DNA_NINJA，或修改 qt-env.cmd 的默认路径。"
    )._p
)

capture(add_heading(doc, "12 个人思考与进一步改进", 1)._p)
capture(add_heading(doc, "12.1 个人想法与更深层次思考：待补充", 2)._p)
capture(
    add_body(
        doc,
        "待补充。此处保留给本人结合课程学习、数据结构取舍、调试过程和展示体验撰写个人思考，"
        "本文档不代写。"
    )._p
)
capture(add_heading(doc, "12.2 客观限制与可验证改进方向", 2)._p)
capture(
    add_body(
        doc,
        "当前只使用 Read 前 K 位作为唯一种子，因此正确位置的前 K 位发生突变时可能没有候选。"
        "可以在不改变核心课程要求的前提下研究多种子或分段种子，但需重新分析候选去重与时间"
        "空间代价。当前随机生成器通过保护前 K 位确保测试 Read 可被该算法召回，这是一项明确"
        "的实验前提。"
    )._p
)
p = add_body(
    doc,
    "当前比对只计算等长字符的汉明距离，不能处理插入、删除、反向互补链和碱基质量分。Bowtie "
    "等成熟工具采用更复杂的压缩索引和回溯策略，本项目不应夸大到真实 NGS 工作负载。若扩展"
    "功能，应先保持现有手写结构的可验证性，再分别引入反向互补查询、多个种子和可选的带状"
    "动态规划。"
)
citation = p.add_run("[4]")
set_run_font(citation, 10.5, superscript=True)
capture(p._p)
capture(
    add_body(
        doc,
        "GUI 当前在主线程中执行批量检索，并通过进度条和事件处理维持反馈。课程规定的数据规模"
        "能够正常完成；若扩大 Reference 或 Reads 数量，应将不可变索引与后台任务结合，同时"
        "设计取消、进度和结果提交边界，避免 UI 模型与核心对象跨线程悬空。"
    )._p
)

capture(add_heading(doc, "13 总结", 1)._p)
capture(
    add_body(
        doc,
        "本课程设计围绕题目 X 完成了可运行的 C++/Qt 桌面程序。核心层手工实现连续 Genome "
        "字符数组、动态桶数组、HashEntry 拉链节点和 K-mer 位置单链表；索引按步长 1 覆盖 "
        "N-K+1 个位置；检索以 Read 前 K 位产生候选，再通过汉明距离和超阈值提前剪枝验证。"
    )._p
)
capture(
    add_body(
        doc,
        "系统将数据生成、文件导入、显式索引构建、单条/批量检索、CSV 导出、全局概览、局部"
        "序列沙盘、精确对齐和 Snackbar 日志反馈组织为完整流程。边界与异常处理围绕错配阈值、"
        "首尾位置、重复 K-mer、非法输入和状态失效进行说明；系统仍受单前缀种子、仅支持替换"
        "错配和课程规模限制，这些边界已在报告中明确。"
    )._p
)

for element in new_blocks:
    move_before(element, references._p)

# References must start on a separate page and participate in the TOC.
references.style = "Heading 1"
set_heading_style(references, 1)
references.paragraph_format.page_break_before = True

# Add an automatic TOC placeholder after the cover and before Chapter 1.
first_heading = next(p for p in doc.paragraphs if p.text.strip().startswith("1 课程设计题目"))
toc_title = first_heading.insert_paragraph_before("目录", style="报告目录标题")
for run in toc_title.runs:
    set_run_font(run, 16, bold=True)
toc_placeholder = first_heading.insert_paragraph_before("[[TOC]]")
toc_placeholder.paragraph_format.first_line_indent = Pt(0)
toc_placeholder.paragraph_format.space_after = Pt(0)
toc_page_break = first_heading.insert_paragraph_before()
toc_page_break.add_run().add_break(WD_BREAK.PAGE)

# Remove a now-stale sentence that pointed to automated tests.
for paragraph in doc.paragraphs:
    if "该边界关系已包含在自动测试中" in paragraph.text:
        new_text = paragraph.text.replace("；该边界关系已包含在自动测试中。", "。")
        paragraph.clear()
        run = paragraph.add_run(new_text)
        set_run_font(run, 10.5)

# Ensure all report text runs use the requested Latin and East Asian fonts.
for paragraph in doc.paragraphs:
    if paragraph.style.name in {
        "报告正文",
        "报告表题",
        "报告参考文献",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "报告目录标题",
    }:
        size = None
        if paragraph.style.name == "Heading 1":
            size = 16
        elif paragraph.style.name == "Heading 2":
            size = 12
        elif paragraph.style.name == "Heading 3":
            size = 14
        elif paragraph.style.name in {"报告正文", "报告表题", "报告参考文献"}:
            size = 10.5
        for run in paragraph.runs:
            set_run_font(run, size)

doc.settings.update_fields_on_open = True
doc.save(OUTPUT)
print(OUTPUT)
