from pathlib import Path
from docx import Document
from lxml import etree
import zipfile
import sys

path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("source.docx")
doc = Document(path)
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip().replace("\t", "<TAB>")
    if text:
        print(f"P{i:03d}\t{p.style.name}\t{text}")
print(f"TABLES\t{len(doc.tables)}")
for i, table in enumerate(doc.tables):
    print(f"T{i}\t{len(table.rows)}x{len(table.columns)}\t{table.cell(0,0).text[:40]}")
with zipfile.ZipFile(path) as z:
    root = etree.fromstring(z.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    fields = root.xpath(".//w:instrText/text()", namespaces=ns)
    print("FIELDS\t" + " | ".join(fields))
