from pathlib import Path
import shutil
import zipfile

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(
    r"D:\Backup\Documents\GitHub\Tongji-CS-2026-spring-my-archive"
    r"\数据结构课程设计"
)
INPUT = ROOT / "tmp" / "report_latest_k_experiment.docx"
TARGET = (
    ROOT
    / "my_works"
    / "250021_张三_计算机技术"
    / "250021_张三_计算机技术_课程设计总结报告.docx"
)


def chapter_text(document: Document) -> str:
    collecting = False
    pieces = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "10 系统运行结果及分析":
            collecting = True
        if collecting:
            pieces.append(text)
        if text == "11 系统安装、运行与操作说明":
            break
    return "\n".join(pieces)


document = Document(INPUT)
text = chapter_text(document)

required_fragments = [
    "建索引耗时为 500 μs",
    "累计检索耗时为 35.90 μs",
    "平均每条为 0.72 μs",
    "建索引耗时由 222 μs 增至 658 μs",
    "批量检索耗时为 466.00 μs",
    "批量检索耗时为 39.30 μs",
    "相对偏移 6～9 中至少存在 1 处突变",
    "外部导入不会根据新 K 改写 Read",
]
for fragment in required_fragments:
    if fragment not in text:
        raise AssertionError(f"Missing required report text: {fragment}")

for stale in ["35.00 μs", "465.70 μs", "33.60 μs", "156 μs", "650 μs"]:
    if stale in text:
        raise AssertionError(f"Stale experiment value remains: {stale}")

tables = [table for table in document.tables if table.cell(0, 0).text.strip() == "K"]
if len(tables) != 1:
    raise AssertionError(f"Expected one K table, got {len(tables)}")

actual_rows = [
    [cell.text.strip() for cell in row.cells]
    for row in tables[0].rows[1:]
]
expected_rows = [
    ["2", "2999", "16", "222", "0.002", "9448 / 466.00 μs / 50"],
    ["6", "2995", "2128", "500", "0.260", "84 / 35.90 μs / 50"],
    ["10", "2991", "2987", "658", "0.365", "37 / 39.30 μs / 37"],
]
if actual_rows != expected_rows:
    raise AssertionError(f"Unexpected experiment table: {actual_rows!r}")

edited_prefixes = (
    "本次实验直接运行",
    "默认参数 K=6",
    "K=6 时索引包含",
    "对比实验保持",
    "随着 K 从 2 增大到 10",
    "候选规模的变化更为明显",
    "K=10 时仅产生",
)
for paragraph in document.paragraphs:
    if not paragraph.text.strip().startswith(edited_prefixes):
        continue
    for run in paragraph.runs:
        run_text = run.text
        if not run_text:
            continue
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        ascii_font = fonts.get(qn("w:ascii"))
        east_asia_font = fonts.get(qn("w:eastAsia"))
        if east_asia_font != "宋体":
            raise AssertionError(f"East Asian font mismatch in run: {run_text!r}")
        contains_cjk = any(
            0x3400 <= ord(char) <= 0x9FFF
            or 0xF900 <= ord(char) <= 0xFAFF
            or 0x3000 <= ord(char) <= 0x303F
            or 0xFF00 <= ord(char) <= 0xFFEF
            for char in run_text
        )
        expected_font = "宋体" if contains_cjk else "Times New Roman"
        if ascii_font != expected_font:
            raise AssertionError(
                f"ASCII font {ascii_font!r} != {expected_font!r} in run {run_text!r}"
            )

with zipfile.ZipFile(INPUT) as archive:
    document_xml = archive.read("word/document.xml")
    if b"\x0b" in document_xml:
        raise AssertionError("Vertical tab found in document XML")
    if b'<w:br/>' in document_xml or b'w:type="textWrapping"' in document_xml:
        raise AssertionError("Soft line break found in document XML")

shutil.copy2(INPUT, TARGET)
print(f"INSTALLED {TARGET}")
print("AUDIT_OK chapter=10.1-10.3 table_rows=3 font_runs=ok soft_breaks=0")
