from pathlib import Path
import shutil

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(
    r"D:\Backup\Documents\GitHub\Tongji-CS-2026-spring-my-archive"
    r"\数据结构课程设计"
)
SOURCE = (
    ROOT
    / "my_works"
    / "250021_张三_计算机技术"
    / "250021_张三_计算机技术_课程设计总结报告.docx"
)
BACKUP = ROOT / "tmp" / "report_before_latest_k_experiment.docx"
OUTPUT = ROOT / "tmp" / "report_latest_k_experiment.docx"


def is_cjk(character: str) -> bool:
    code = ord(character)
    return (
        0x3400 <= code <= 0x4DBF
        or 0x4E00 <= code <= 0x9FFF
        or 0xF900 <= code <= 0xFAFF
        or 0x3000 <= code <= 0x303F
        or 0xFF00 <= code <= 0xFFEF
    )


def replace_paragraph(paragraph, text: str) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)

    if not text:
        return

    start = 0
    current_is_cjk = is_cjk(text[0])
    for index in range(1, len(text) + 1):
        boundary = index == len(text)
        next_is_cjk = None if boundary else is_cjk(text[index])
        if boundary or next_is_cjk != current_is_cjk:
            run = paragraph.add_run(text[start:index])
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:eastAsia"), "宋体")
            if current_is_cjk:
                fonts.set(qn("w:ascii"), "宋体")
                fonts.set(qn("w:hAnsi"), "宋体")
                run.font.name = "宋体"
            else:
                fonts.set(qn("w:ascii"), "Times New Roman")
                fonts.set(qn("w:hAnsi"), "Times New Roman")
                run.font.name = "Times New Roman"
            if not boundary:
                start = index
                current_is_cjk = next_is_cjk


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


if not SOURCE.exists():
    raise FileNotFoundError(SOURCE)

shutil.copy2(SOURCE, BACKUP)
document = Document(SOURCE)

replacements = {
    "本次实验直接运行": (
        "本次实验直接运行提交目录中的 Release 版 DNA_SearchEngine.exe，导入配套的 "
        "test_reference.fasta 和 test_reads.txt。Reference 长度为 3000 bp，共导入 50 条"
        "长度为 30 bp 的 Reads；允许错配数固定为 k=2，只改变 K-mer 长度 K。对 K=2、6、10 "
        "分别显式建立索引并执行一次高通量批量检索，随后导出 DNA_Search_Results_K2.csv、"
        "DNA_Search_Results_K6.csv 和 DNA_Search_Results_K10.csv。不同 K-mer 数、建索引耗时"
        "和哈希负载取自同次运行界面，候选数、成功数、汉明距离及检索耗时由相应 CSV 逐行统计。"
        "三次实验使用同一组导入 Reads，而不是在改变 K 后重新随机生成，因此可以直接比较索引"
        "参数变化对候选规模和检索结果的影响。"
    ),
    "默认参数 K=6": (
        "默认参数 K=6、k=2 时，50 条 Reads 全部检索成功，成功数为 50，失败数为 0。CSV 共 "
        "50 条结果记录，每条 Read 恰好对应 1 个合法匹配位置，没有重复或多位置结果。汉明距离"
        "分布为 d=1 的 25 条、d=2 的 25 条，与测试数据中交替设置 1～2 个替换突变的生成规则"
        "一致。每条结果的错配详情均给出了参考基因组绝对位置和碱基替换关系，可与界面的精确"
        "比对详情逐项对应。"
    ),
    "K=6 时共验证": (
        "K=6 时索引包含 2128 种不同 K-mer，建索引耗时为 500 μs，哈希负载为 0.260。检索"
        "阶段共验证 84 个种子候选，平均每条 Read 为 1.68 个，单条候选数范围为 1～5。50 条 "
        "Reads 的累计检索耗时为 35.90 μs，平均每条为 0.72 μs，单条耗时范围为 0.40～1.30 μs。"
        "导入的测试 Reads 在生成时按 K=6 保护前缀种子，因此正确位置均能进入候选集合；后续"
        "完整 Read 验证准确保留了 1 处或 2 处错配的结果，说明“种子筛选—汉明距离验证—超阈值"
        "剪枝”的流程在默认参数下与预期一致。"
    ),
    "对比实验保持": (
        "对比实验保持 Reference、Reads、Read 长度和允许错配数不变，依次将 K 设置为 2、6、10。"
        "每次修改 K 后旧索引自动失效，再重新点击“建立索引”，最后执行同一批 50 条 Reads 的"
        "批量检索。长度为 N=3000 的 Reference 在三种 K 下分别处理 N-K+1，即 2999、2995、"
        "2991 个滑动窗口；索引和检索统计见表 7。"
    ),
    "随着 K 从 2 增大到 10": (
        "随着 K 从 2 增大到 10，不同 K-mer 数由 16 增至 2987，哈希负载由 0.002 增至 0.365，"
        "建索引耗时由 222 μs 增至 658 μs。K=2 只有 4²=16 种可能短词，大量滑动窗口被归并到"
        "少数键的位置链表；K 增大后短词区分度提高，独立键数量显著增加，需要创建和插入更多"
        "哈希节点，因此本次实测的建索引耗时随之上升。"
    ),
    "候选规模的变化更为明显": (
        "候选规模的变化更为明显：K=2 平均每条 Read 产生 188.96 个候选，累计 9448 个候选，"
        "批量检索耗时为 466.00 μs；K=6 仅产生 84 个候选，累计耗时降至 35.90 μs。两种参数"
        "均成功检索 50 条 Reads，且对应匹配位置完全一致。由此可见，在本组数据上 K=6 在保持"
        "召回结果的同时大幅减少了完整 Read 的逐字符验证次数，并取得三组中最低的实测批量"
        "检索耗时，是较合适的取值。"
    ),
    "K=10 时仅产生": (
        "K=10 时仅产生 37 个候选，批量检索耗时为 39.30 μs，成功数为 37。逐行核对发现，未命中"
        "的 Read 编号为 1、8、15、16、18、22、25、32、39、40、42、46、49；这 13 条 Read 在"
        "相对偏移 6～9 中至少存在 1 处突变，因而突变碱基进入 K=10 的前缀种子。test_reads.txt "
        "是按默认 K=6 生成后保存并在本实验中导入的固定数据，只保证前 6 位不突变；外部导入"
        "不会根据新 K 改写 Read，所以该现象符合单一精确前缀种子的检索机制，并非汉明距离验证"
        "错误。K=10 已命中的 37 条 Read 与 K=6 的匹配位置一致。其候选数虽最少，但本次耗时比 "
        "K=6 高 3.40 μs；在几十微秒量级的单次测量中不宜据此作过度性能推断。若要公平评价 "
        "K=10 的随机数据召回率，应在 K=10 下重新生成动态保护前 10 位的 Reads，或在后续改进"
        "中采用多种子策略。"
    ),
}

for prefix, replacement in replacements.items():
    replace_paragraph(find_paragraph(document, prefix), replacement)

experiment_tables = [
    table for table in document.tables if table.cell(0, 0).text.strip() == "K"
]
if len(experiment_tables) != 1:
    raise RuntimeError(f"Expected one K experiment table, found {len(experiment_tables)}")

table = experiment_tables[0]
rows = [
    ["2", "2999", "16", "222", "0.002", "9448 / 466.00 μs / 50"],
    ["6", "2995", "2128", "500", "0.260", "84 / 35.90 μs / 50"],
    ["10", "2991", "2987", "658", "0.365", "37 / 39.30 μs / 37"],
]
for row_index, values in enumerate(rows, start=1):
    for column_index, value in enumerate(values):
        cell = table.cell(row_index, column_index)
        paragraphs = cell.paragraphs
        replace_paragraph(paragraphs[0], value)
        for extra in paragraphs[1:]:
            replace_paragraph(extra, "")

document.save(OUTPUT)
print(OUTPUT)
