from pathlib import Path

from docx import Document

path = Path(
    r"my_works\250021_张三_计算机技术"
    r"\250021_张三_计算机技术_课程设计总结报告.docx"
)
document = Document(path)

in_chapter = False
for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if "系统运行结果及分析" in text and paragraph.style.name.startswith("Heading"):
        in_chapter = True
    if in_chapter:
        print(index, paragraph.style.name, text.encode("unicode_escape"))
    if in_chapter and "系统安装、运行与操作说明" in text:
        break

for table_index, table in enumerate(document.tables):
    first = table.cell(0, 0).text.strip()
    if first in {"记录项", "K"}:
        print(f"TABLE {table_index}")
        for row in table.rows:
            print([cell.text for cell in row.cells])
